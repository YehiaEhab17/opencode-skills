#!/usr/bin/env python3
"""Run this to verify the parsing-documents skill is working correctly."""

import subprocess
import sys
import os

PASS = 0
FAIL = 0


def check(name, ok, detail=""):
    global PASS, FAIL
    status = "PASS" if ok else "FAIL"
    if ok:
        PASS += 1
    else:
        FAIL += 1
    print(f"  [{status}] {name}" + (f" — {detail}" if detail else ""))


print("=== Parsing Documents Setup Test ===\n")

# 1. markitdown CLI
print("1. MarkItDown CLI")
try:
    r = subprocess.run(
        ["markitdown", "--help"], capture_output=True, text=True, timeout=10
    )
    check("markitdown CLI exists", r.returncode == 0)
except FileNotFoundError:
    check(
        "markitdown CLI exists", False, "not on PATH — run: pip install markitdown[all]"
    )

# 2. Plugins
print("\n2. Plugins")
r = subprocess.run(
    ["markitdown", "--list-plugins"], capture_output=True, text=True, timeout=10
)
check(
    "markitdown-math plugin found",
    "markitdown-math" in r.stdout,
    "run: pip install -e scripts/markitdown-math",
)
check(
    "markitdown-ocr plugin found", "ocr" in r.stdout, "run: pip install markitdown-ocr"
)

# 3. Python imports
print("\n3. Python imports")
errors = []
for mod in ["markitdown", "fitz", "pdfplumber", "pdfminer", "pptx", "docx"]:
    try:
        __import__(mod)
    except ImportError:
        errors.append(mod)
check(
    "All core modules importable",
    not errors,
    f"missing: {', '.join(errors)}" if errors else "",
)

# 4. Plugin converter registration
print("\n4. Plugin converters")
try:
    from markitdown import MarkItDown

    md = MarkItDown(enable_plugins=True)
    names = [reg.converter.__class__.__name__ for reg in md._converters]
    check("MathPptxConverter registered", "MathPptxConverter" in names)
    check("EnhancedPdfConverter registered", "EnhancedPdfConverter" in names)
    check(
        "Converters at priority -2.0",
        any(reg.priority == -2.0 for reg in md._converters),
    )
except Exception as e:
    check("Plugin registration", False, str(e))

# 5. File conversion smoke tests
print("\n5. File conversion smoke tests")
from markitdown import MarkItDown
from markitdown._stream_info import StreamInfo
import io

md = MarkItDown(enable_plugins=True)

# PPTX acceptance
pptx_info = StreamInfo(
    extension=".pptx",
    mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
)
pptx_ok = any(
    reg.converter.accepts(io.BytesIO(b"fake"), pptx_info)
    for reg in md._converters
    if "MathPptx" in reg.converter.__class__.__name__
)
check("MathPptxConverter accepts .pptx", pptx_ok)

# PDF acceptance
pdf_info = StreamInfo(extension=".pdf", mimetype="application/pdf")
pdf_ok = any(
    reg.converter.accepts(io.BytesIO(b"fake"), pdf_info)
    for reg in md._converters
    if "EnhancedPdf" in reg.converter.__class__.__name__
)
check("EnhancedPdfConverter accepts .pdf", pdf_ok)

# Real DOCX conversion
from docx import Document

doc = Document()
doc.add_heading("Real DOCX Test", 1)
doc.add_paragraph("This is a test paragraph.")
doc.save("/tmp/test_setup.docx")
try:
    r = subprocess.run(
        ["markitdown", "/tmp/test_setup.docx"],
        capture_output=True,
        text=True,
        timeout=15,
    )
    check("Real DOCX converts successfully", "Real DOCX Test" in r.stdout)
except Exception as e:
    check("Real DOCX converts successfully", False, str(e))

# 6. Ollama check (optional)
print("\n6. Ollama vision model (optional)")
try:
    import requests

    r = requests.get("http://localhost:11434/api/tags", timeout=3)
    if r.status_code == 200:
        models = r.json().get("models", [])
        vision_models = [
            m["name"]
            for m in models
            if any(
                v in m["name"].lower()
                for v in [
                    "llava",
                    "moondream",
                    "minicpm",
                    "bakllava",
                    "vision",
                    "gemma3",
                ]
            )
        ]
        if vision_models:
            check(
                "Ollama running with vision model",
                True,
                f"found: {', '.join(vision_models)}",
            )
        elif models:
            check(
                "Ollama running (no vision model)",
                True,
                f"models: {', '.join(m['name'] for m in models)} — install one: ollama pull llava",
            )
        else:
            check(
                "Ollama running (no models)",
                True,
                "no models pulled — run: ollama pull llava",
            )
    else:
        check("Ollama running", False, f"unexpected status: {r.status_code}")
except requests.ConnectionError:
    check(
        "Ollama running",
        False,
        "not found at localhost:11434 — start it with: ollama serve",
    )
except ImportError:
    check("Ollama check skipped", True, "requests not installed")

# 7. Legacy script check
print("\n7. Legacy converter")
legacy_path = os.path.join(os.path.dirname(__file__), "convert_legacy.sh")
check("convert_legacy.sh exists", os.path.isfile(legacy_path))
check("convert_legacy.sh is executable", os.access(legacy_path, os.X_OK))

# Summary
print(f"\n{'=' * 40}")
print(f"Results: {PASS} passed, {FAIL} failed, {PASS + FAIL} total")
if FAIL == 0:
    print("All checks passed — setup is ready to use!")
else:
    print(f"Fix {FAIL} failure(s) above before using the skill.")
    sys.exit(1)
